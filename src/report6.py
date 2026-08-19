"""v6 production report (PLAN_V6): second-expansion features, gated by F4/T1.

v6 = v5 pipeline with F4 gate survivors (per params_v6); T1 travel terms would
enter d_adj if they had shipped (they are recorded in params_v6 either way).
Everything else inherits from params_v4/v5 shipped. Per PLAN_V6 H: live 2026-27
holdout still scores v1/v4/HOWE; v6 + HOWE6 = 0.5*v1 + 0.5*v6 co-headline
2027-28; 2026-27 v6 numbers are report-only. Seeds 611 (h1) / 622 (h2), offsets
v6=0 / HOWE6=1. Writes output/nhl_projections_v6.xlsx + CSVs + v6_prior_ratings.
"""
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import engine as E
import goalie_game as GG
import players as P
import ridge as R
from features_v6 import FEATURES_V6_ALL, FeatureBuilderV6
from report import FULL
from report4 import prod_noise, real_schedule_flags, summarize
from report5 import BLOCKS_V5

PROJ = Path(__file__).resolve().parents[1]
RAW = PROJ / "data" / "raw"
OUT = PROJ / "output"
N_SIMS = 10_000
V_PROD = 2026

BLOCKS_V6 = dict(BLOCKS_V5)
BLOCKS_V6["Contrib Structure"] = ["line_cont", "toi_hhi_f", "coach_new",
                                  "coach_tenure"]
BLOCKS_V6["Contrib Prospects"] = ["prospect_pipeline", "prospect_prod"]


def production_predictions_v6(fb, feats, lam, h, target_T):
    pred_seasons = list(range(2012 + (h - 1), 2027)) + [target_T]
    X, y, meta = fb.feature_matrix(pred_seasons, h, feats=FEATURES_V6_ALL)
    idx = [FEATURES_V6_ALL.index(f) for f in feats]
    Xs = X[:, idx]
    tr = ~np.isnan(y)
    b = R.fit_ridge(Xs[tr], y[tr], lam)
    te = meta["T"].to_numpy() == target_T
    yhat = Xs[te] @ b
    yhat -= yhat.mean()
    teams = meta.loc[te, "team"].to_numpy()
    contribs = {}
    for label, fs in BLOCKS_V6.items():
        sub = [f for f in fs if f in feats]
        if not sub:
            continue
        ii = [feats.index(f) for f in sub]
        cv = Xs[te][:, ii] @ b[ii]
        contribs[label] = pd.Series(cv - cv.mean(), index=teams)
    return pd.Series(yhat, index=teams), contribs


def main():
    p2 = json.loads((OUT / "params_v2.json").read_text())
    p4 = json.loads((OUT / "params_v4.json").read_text())
    p6 = json.loads((OUT / "params_v6.json").read_text())
    v1p = json.loads((OUT / "params.json").read_text())
    ship4, ship6 = p4["shipped"], p6["shipped"]
    c, k = p2["c"], p2["k"]
    g, ts = E.load()
    preds, end_r, _ = E.run_elo(g, K=v1p["K"], H=v1p["H"], phi_s=v1p["phi_s"])
    fb = FeatureBuilderV6(end_r, ts, goalie_hp=p2["goalie_hp"],
                          skater_delta=p2["skater_delta"])
    om = E.fit_outcome(preds, list(range(2006, 2027)))
    beta_prod = E.fit_xg_beta(end_r, ts, list(range(2008, 2027)))
    sk, go, skt, got, bios = P.load_panels()
    delo = pd.read_csv(OUT / "overlay_team_deltas_v4.csv", index_col=0)["dElo"]
    sigma = ship4["sigma_c4"]

    results, prior_rows = {}, {}
    for h, (label, target_T, seed) in {1: ("2026_27", 2027, 611),
                                       2: ("2027_28", 2028, 622)}.items():
        feats = ship6[f"features_h{h}"]
        lam = ship6[f"lam_h{h}"]
        yhat, contribs = production_predictions_v6(fb, feats, lam, h, target_T)
        ov_scale = 1.0 if h == 1 else 0.88
        r_v6 = {t: 1505.0 + yhat[t] / c + ov_scale * float(delo.get(t, 0.0))
                for t in yhat.index}
        phi = v1p["phi1"] if h == 1 else v1p["phi2"]
        r_v1 = E.project_ratings(end_r, ts, beta_prod, V_PROD, w=v1p["w"], phi=phi)
        r_howe6 = {t: 1505.0 + 0.5 * (r_v1[t] - 1505.0) + 0.5 * (r_v6[t] - 1505.0)
                   for t in r_v6}

        noise_fn, nteams, par2, gpar, tand, frag = prod_noise(
            fb, sk, got, bios, ts, k, h, ship4["n0_g"])
        rng = np.random.default_rng(seed)
        if h == 1:
            sched = real_schedule_flags()
        else:
            sched = E.synthetic_schedule_84(E.DIVISIONS_CURRENT, rng)
            sched["hb2b"] = 0.0
            sched["ab2b"] = 0.0
            sched["d_adj"] = 0.0
        gn = (GG.make_game_noise(sched, tand, k) if ship4["goalie_layer"] else None)
        sim_cols = ["home", "away", "d_adj"]
        for name, ratings, seed_off in (("v6", r_v6, 0), ("HOWE6", r_howe6, 1)):
            assert sorted(ratings) == sorted(nteams)
            sim = E.simulate_season(ratings, sigma, sched[sim_cols], om,
                                    E.DIVISIONS_CURRENT, N_SIMS,
                                    np.random.default_rng(seed + seed_off),
                                    playoffs=True, extra_noise=noise_fn, game_noise=gn)
            assert sim["made_po"].sum(1).std() == 0 and sim["won_cup"].sum(1).mean() == 1.0
            extra = {"Avail SD (Elo)": frag.avail_sd_elo,
                     "Overlay dElo": pd.Series({t: round(ov_scale * float(delo.get(t, 0.0)), 1)
                                                for t in ratings})}
            if name == "v6":
                for lbl in contribs:
                    extra[lbl] = contribs[lbl]
            results[(label, name)] = summarize(sim, extra)
            print(f"{label} {name}: league mean {sim['pts'].mean():.1f}; top "
                  f"{results[(label, name)].iloc[0].Team} "
                  f"{results[(label, name)].iloc[0].xPts:.1f}")
        if h == 1:
            prior_rows = {t: {"rating_v1": round(r_v1[t], 2),
                              "rating_v6": round(r_v6[t], 2),
                              "rating_howe6": round(r_howe6[t], 2)}
                          for t in sorted(r_v6)}

    pd.DataFrame(prior_rows).T.rename_axis("team").to_csv(OUT / "v6_prior_ratings.csv")

    readme = [
        "NHL 2026-27 / 2027-28 — MODEL v6 (built 2026-08-19; PLAN_V6 prereg @ b719257)",
        "",
        "v6 = v5 + second-expansion features gated by F4 (train-only 2012-2017):",
        f"h1 adds {p6['gate_F4']['h1']['joint']['new']} "
        f"(LOSO {p6['gate_F4']['h1']['mae_incumbent']} -> "
        f"{p6['gate_F4']['h1']['joint']['mae_final']}), "
        f"h2 adds {p6['gate_F4']['h2']['joint']['new']} "
        f"(LOSO {p6['gate_F4']['h2']['mae_incumbent']} -> "
        f"{p6['gate_F4']['h2']['joint']['mae_final']}).",
        f"Travel gate T1: {p6['gate_T1']['ships'] or 'NULL (documented)'}. "
        f"S3 spell screen: EB beats bucket = {p6['screen_S3']['eb_beats_bucket']} "
        "(recorded; availability rebuild deferred).",
        "Data: shift charts (pair TOI), coach records, prospect production "
        "(records.nhl.com + player-landing careers), travel table, playoff PBP, "
        "EDGE player tracking, odds log.",
        "",
        "IMPORTANT (PLAN_V6 H): live 2026-27 holdout still scores v1/v4/HOWE. The "
        "2026-27 v6 sheets are REPORT-ONLY. v6 and HOWE6 = 0.5*v1 + 0.5*v6 are the "
        "2027-28 co-headline.",
    ]
    xlsx = OUT / "nhl_projections_v6.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        pd.DataFrame({"NHL Model v6": readme}).to_excel(xw, sheet_name="README", index=False)
        results[("2027_28", "HOWE6")].to_excel(xw, sheet_name="Projections_2027_28_HOWE6", index=False)
        results[("2027_28", "v6")].to_excel(xw, sheet_name="Projections_2027_28_v6", index=False)
        results[("2026_27", "HOWE6")].to_excel(xw, sheet_name="ReportOnly_2026_27_HOWE6", index=False)
        results[("2026_27", "v6")].to_excel(xw, sheet_name="ReportOnly_2026_27_v6", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="4A1D6E")
                cell.font = Font(bold=True, color="FFFFFF")
    results[("2026_27", "v6")].to_csv(OUT / "projections_2026_27_v6.csv", index=False)
    results[("2027_28", "v6")].to_csv(OUT / "projections_2027_28_v6.csv", index=False)
    results[("2026_27", "HOWE6")].to_csv(OUT / "projections_2026_27_howe6.csv", index=False)
    results[("2027_28", "HOWE6")].to_csv(OUT / "projections_2027_28_howe6.csv", index=False)
    print(f"wrote {xlsx}")


if __name__ == "__main__":
    main()
