"""v2 production report: 2026-27 headline (real schedule) + 2027-28 bonus.
Writes output/nhl_2026_27_projections_v2.xlsx + CSVs. v1 outputs untouched (verified)."""
import json
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import engine as E
import players as P
import ridge as R
from features import FeatureBuilder, FEATURES
from report import FULL, fair_odds

PROJ = Path(__file__).resolve().parents[1]
RAW = PROJ / "data" / "raw"
OUT = PROJ / "output"
N_SIMS = 10_000
V_PROD = 2026

BLOCKS = {
    "Contrib Elo/xG": ["elo_dev", "xg_dev"],
    "Contrib Goaltending": ["gsax_1yr", "gsax_marcel", "tandem_gsax", "goalie_consistency",
                            "goalie_age", "goalie_trend"],
    "Contrib Age": ["toi_age", "share_u23", "share_32p", "prod_age_exp"],
    "Contrib Roster": ["returning_toi", "star_share", "draft_cap", "player_points_proj"],
}


def real_schedule() -> pd.DataFrame:
    df = pd.read_csv(RAW / "nhl_schedule_20262027.csv")[["home", "away"]]
    return df


def production_predictions(fb, p2, h, target_T):
    lam = p2["lams"][f"h{h}"]["full"]
    pred_seasons = list(range(2012 + (h - 1), 2027)) + [target_T]
    X, y, meta = fb.feature_matrix(pred_seasons, h)
    seasons = meta["T"].to_numpy()
    tr = ~np.isnan(y)
    b = R.fit_ridge(X[tr], y[tr], lam)
    te = seasons == target_T
    yhat = X[te] @ b
    yhat = yhat - yhat.mean()
    teams = meta.loc[te, "team"].to_numpy()
    contribs = {}
    for label, feats in BLOCKS.items():
        idx = [FEATURES.index(f) for f in feats]
        cv = X[te][:, idx] @ b[idx]
        contribs[label] = pd.Series(cv - cv.mean(), index=teams)
    return pd.Series(yhat, index=teams), contribs, b


def summarize_v2(sim, yhat, contribs, delo, sigma, c):
    teams = sim["teams"]
    pts = sim["pts"].astype(float)
    div_of = {t: dv for dv, ts_ in E.DIVISIONS_CURRENT.items() for t in ts_}
    conf_of = {t: cf for cf, dvs in E.CONFS.items() for dv in dvs
               for t in E.DIVISIONS_CURRENT[dv]}
    key = pts * 1e8 + sim["rw"] * 1e4 + sim["row"] + np.random.default_rng(1).random(pts.shape)
    top_idx = key.argmax(axis=1)
    rows = []
    for i, t in enumerate(teams):
        p = pts[:, i]
        playoff = sim["made_po"][:, i].mean()
        cup = sim["won_cup"][:, i].mean()
        po_dec, _ = fair_odds(playoff)
        cup_dec, cup_amer = fair_odds(cup)
        row = {
            "Team": FULL[t], "Abbr": t, "Conference": conf_of[t], "Division": div_of[t],
            "xPts": p.mean(), "SD": p.std(ddof=1),
            "P5": np.percentile(p, 5), "P25": np.percentile(p, 25),
            "P50": np.percentile(p, 50), "P75": np.percentile(p, 75),
            "P95": np.percentile(p, 95),
            "Playoff%": playoff, "Division%": sim["won_div"][:, i].mean(),
            "Presidents%": float((top_idx == i).mean()),
            "Conference%": sim["won_conf"][:, i].mean(), "Cup%": cup,
            "P(>=100pts)": (p >= 100).mean(),
            "Ridge dev (pts/82)": float(yhat.get(t, 0.0)),
        }
        for label in BLOCKS:
            row[label] = float(contribs[label].get(t, 0.0))
        row["Overlay dElo"] = float(delo.get(t, 0.0))
        row["Sigma (Elo)"] = float(sigma if np.isscalar(sigma) else sigma[t])
        row["FairOdds Playoff (dec)"] = po_dec
        row["FairOdds Cup (dec)"] = cup_dec
        row["FairOdds Cup (US)"] = cup_amer
        rows.append(row)
    df = pd.DataFrame(rows).sort_values("xPts", ascending=False).reset_index(drop=True)
    df.insert(0, "Rank", np.arange(1, len(df) + 1))
    return df


def goalie_sheet(fb, bios, go, ros):
    proj = fb.goalie_proj(V_PROD).copy()
    names = go.drop_duplicates("playerId").set_index("playerId")["name"]
    api_team = ros[ros.position == "G"].drop_duplicates("playerId").set_index("playerId").team
    proj["Name"] = proj.playerId.map(names)
    proj["2026-27 team"] = proj.playerId.map(api_team)
    proj["Age"] = P.age_of(bios, proj.playerId, 2027).round(1).to_numpy()
    proj["GSAx/1000 shots (proj)"] = (proj.theta * 1000).round(2)
    proj["GSAx/season (starter)"] = (proj.theta * 2500).round(1)
    proj["Posterior SD/1000"] = (np.sqrt(proj.V_post) * 1000).round(2)
    proj["Consistency C"] = proj.C.round(2)
    proj["Steadiness pctile"] = (100 * (1 - proj.C.rank(pct=True))).round(0)
    proj["Trend/1000/yr"] = (proj.trend * 1000).round(2)
    keep = proj[proj["2026-27 team"].notna() | (proj.shots_win >= 2000)]
    keep = keep.sort_values("theta", ascending=False)
    return keep[["Name", "2026-27 team", "Age", "GSAx/1000 shots (proj)",
                 "GSAx/season (starter)", "Posterior SD/1000", "Consistency C",
                 "Steadiness pctile", "Trend/1000/yr", "m", "shots_win"]].rename(
        columns={"m": "Seasons used", "shots_win": "Shots (window)"})


def age_sheet(fb, sk, bios):
    a = fb.team_age_structure(V_PROD, 1).round(3).reset_index()
    a["Team"] = a.team.map(FULL)
    fe = fb.fe_coefs(V_PROD)
    ages = np.arange(19, 41)
    curves = pd.DataFrame({"age": ages})
    for pg in ("F", "D"):
        c = P.curve_value(fe[pg], ages)
        curves[f"{pg} pts60 vs peak"] = (c - c.max()).round(3)
    return a, curves


def player_points_sheet(fb, sk, bios, ros):
    m = fb.marcel(V_PROD, 1).set_index("playerId")
    names = sk.drop_duplicates("playerId").set_index("playerId")["name"]
    r = ros[ros.position != "G"].drop_duplicates("playerId").copy()
    r["Name"] = r.playerId.map(names)
    r = r[r.playerId.isin(m.index)]
    r["Proj Pts/60"] = r.playerId.map(m.proj_pts60).round(2)
    r["Proj TOI/82"] = r.playerId.map(m.toi82_proj).round(0)
    r["Age adj"] = r.playerId.map(m.age_adj).round(3)
    r["Proj Points"] = r.playerId.map(m.proj_points).round(1)
    r["Age"] = P.age_of(bios, r.playerId, 2027).round(1).to_numpy()
    rows = []
    for team, d in r.groupby("team"):
        d = d.sort_values("Proj Points", ascending=False).head(12)
        for rank, (_, x) in enumerate(d.iterrows(), 1):
            rows.append({"Team": FULL.get(team, team), "Rk": rank, "Player": x["Name"],
                         "Pos": x.position, "Age": x.Age, "Proj Points": x["Proj Points"],
                         "Proj Pts/60": x["Proj Pts/60"], "Proj TOI/82": x["Proj TOI/82"],
                         "Age adj": x["Age adj"]})
    return pd.DataFrame(rows)


def backtest_sheet(p2):
    rows = [{"section": "PRE-REGISTRATION", "detail": json.dumps(p2["prereg"], indent=0)}]
    rows.append({"section": "train (LOSO MAE, <=2017)", "detail": json.dumps(p2["train_mae"])})
    rows.append({"section": "confirm gate (2018-2021, h1 excl 2021)",
                 "detail": json.dumps(p2["confirm_result"]["gate"])
                 + f" -> chosen {p2['confirm_result']['choice']}"})
    for h in (1, 2):
        for rec in p2["final_result"][f"h{h}"]:
            rows.append({"section": f"final h{h}", "detail": json.dumps(rec)})
    return pd.DataFrame(rows)


def readme_lines(p2, agg):
    L = []
    a = L.append
    fr = {h: pd.DataFrame(p2["final_result"][f"h{h}"]) for h in (1, 2)}
    a("NHL 2026-27 PROJECTIONS - MODEL v2 (pieces-driven) - built 2026-08-13, data through 2025-26")
    a("")
    a("v2 = v1's validated Elo+xG backbone, generalized to a ridge regression over 16 team features")
    a("built from the PIECES: goaltending (empirical-Bayes GSAx projections with per-goalie")
    a("CONSISTENCY ratings), age structure (player-level TOI-weighted curves, peak ~27, D flatter),")
    a("points-by-players (Marcel-projected rosters), organizational churn, star concentration, and")
    a("draft capital. Feature Ledger records every candidate with its measured effect - nothing was")
    a("rejected without a number.")
    a("")
    a("HEADLINE: Projections_2026_27 uses the REAL published 84-game 2026-27 schedule (42H/42A).")
    a("Projections_2027_28 is a two-season-ahead bonus (synthetic CBA schedule).")
    a("")
    a("=== Validation (three-stage, honest about the spent window) ===")
    a("Train <=2017 tuned everything; Confirm 2018-2021 selected the model by pre-registered gate;")
    a("Final 2022-2026 was run ONCE with no decisions. The 2018-2026 window was previously used")
    a("once to validate v1; v2 re-partitioned it (confirm/final) to quarantine final-window seasons")
    a("from all v2 choices.")
    a(f"Confirm (h1 dev-MAE excl 2021): R3 {p2['confirm_result']['gate']['R3'][0]} vs R2 "
      f"{p2['confirm_result']['gate']['R2'][0]} vs v1 {p2['confirm_result']['gate']['R1'][0]} "
      f"-> R3 (full pieces model) adopted")
    a(f"Final h1 mean dev-MAE: v2 {fr[1].mae_v2.mean():.2f} | v1 {fr[1].mae_v1.mean():.2f} | "
      f"regressed-prior {fr[1].mae_regressed.mean():.2f} | uniform {fr[1].mae_uniform.mean():.2f}")
    a(f"Final h1 mean Spearman: v2 {fr[1].spearman_v2.mean():.3f} | v1 {fr[1].spearman_v1.mean():.3f}")
    a(f"Final h2 mean dev-MAE: v2 {fr[2].mae_v2.mean():.2f} | v1 {fr[2].mae_v1.mean():.2f} | "
      f"regressed-prior {fr[2].mae_regressed.mean():.2f}")
    a("")
    a("HONEST VERDICT: v2 beat v1 on the confirm window (-0.38 MAE) and lost the final window")
    a("(+0.21); with 4-5 seasons per window (paired SE ~0.35) the two are statistically")
    a("indistinguishable on accuracy. What v2 adds at no accuracy cost: per-team attribution")
    a("(Elo/xG vs goaltending vs age vs roster columns), goalie consistency ratings, player-level")
    a("projections, and tested-and-documented nulls. The rung choice was made at confirm per the")
    a("pre-registration and is NOT revisited in light of the final window - that discipline is")
    a("the point. 2025-26 note: in extreme parity seasons the uniform baseline beats every skill")
    a("model; expect compressed spreads in 2026-27 too.")
    a("")
    a("=== What the pieces contributed (see Feature_Ledger) ===")
    a("- Goaltending: projected tandem GSAx + team GSAx history carry real weight; goalie talent")
    a("  stabilizes around ~5,000 shots; consistency C separates steady multi-year goalies (C<1,")
    a("  e.g. Hellebuyck/Shesterkin ~0.6) from one-great-season profiles (C>1).")
    a("- Age: 'old teams regress, young teams step forward' shows up mainly at the two-season")
    a("  horizon; production-age exposure and 32+ TOI share are the carriers.")
    a("- NULLS kept honest: trajectory/momentum (exactly 0), goalie-uncertainty heteroscedastic")
    a("  sims (pre-registered coverage test: no improvement -> dropped), and the July-2026")
    a("  roster-delta overlay: mechanism real (rho=0.74, t~3.3) but the placebo test failed")
    a("  (-0.53): teams that add value were last season's underperformers (selection), so the")
    a("  pre-registered rule BLOCKED application. Movers are valued in Roster_Delta_2026 but add")
    a("  0 Elo. This is the protocol working, not a missing feature.")
    a("")
    a("=== League aggregates (2026-27, 10,000 sims) ===")
    for k_, v_ in agg.items():
        a(f"{k_}: {v_}")
    a("")
    a("=== Caveats ===")
    a("- August-2026 rosters are ~25-man snapshots; late signings/trades/PTOs not reflected.")
    a("- Player projections are two-way-uncertain: Marcel + age curve, no injuries/deployment.")
    a("- Sources: Hockey-Reference (results), MoneyPuck.com (xG, player data; non-commercial")
    a("  attribution), NHL API (rosters/draft/schedule). Fair odds are break-even, no vig.")
    return L


def main():
    p2 = json.loads((OUT / "params_v2.json").read_text())
    assert "final_result" in p2, "run backtest2.py final first"
    choice = p2["confirm_result"]["choice"]
    assert choice == "R3", f"report2 written for R3; got {choice}"
    v1 = json.loads((OUT / "params.json").read_text())
    g, ts = E.load()
    preds, end_r, _ = E.run_elo(g, K=v1["K"], H=v1["H"], phi_s=v1["phi_s"])
    fb = FeatureBuilder(end_r, ts, goalie_hp=p2["goalie_hp"], skater_delta=p2["skater_delta"])
    om = E.fit_outcome(preds, list(range(2006, 2027)))
    c = p2["c"]
    sk, go, skt, got, bios = P.load_panels()
    ros = pd.read_csv(RAW / "nhl_rosters_20262027.csv")
    delo = pd.read_csv(OUT / "overlay_team_deltas.csv", index_col=0)["dElo"]

    results, sims, aggs = {}, {}, {}
    for h, (label, target_T, seed) in {1: ("2026_27", 2027, 111),
                                       2: ("2027_28", 2028, 222)}.items():
        yhat, contribs, beta = production_predictions(fb, p2, h, target_T)
        ratings = {t: 1505.0 + yhat[t] / c + (delo.get(t, 0.0) if h == 1 else 0.0)
                   for t in yhat.index}
        rng = np.random.default_rng(seed)
        sched = real_schedule() if h == 1 else E.synthetic_schedule_84(E.DIVISIONS_CURRENT, rng)
        assert set(sched.home) == set(ratings), "schedule/team mismatch"
        sigma = p2["sigma_c"][f"h{h}"]
        sim = E.simulate_season(ratings, sigma, sched, om, E.DIVISIONS_CURRENT,
                                N_SIMS, rng, playoffs=True)
        assert sim["made_po"].sum(1).mean() == 16 and sim["made_po"].sum(1).std() == 0
        assert sim["won_cup"].sum(1).mean() == 1.0
        results[label] = summarize_v2(sim, yhat, contribs, delo if h == 1 else pd.Series(dtype=float),
                                      sigma, c)
        sims[label] = sim
        pts = sim["pts"].astype(float)
        aggs[label] = {
            "league mean points": round(float(pts.mean()), 2),
            "mean SD of team points": round(float(pts.std(axis=1, ddof=1).mean()), 2),
            "expected #teams >=100": round(float((pts >= 100).sum(1).mean()), 2),
            "expected Presidents' points": round(float(pts.max(axis=1).mean()), 1),
            "expected last-place points": round(float(pts.min(axis=1).mean()), 1),
        }
        print(f"{label}: league mean {pts.mean():.1f}; top {results[label].iloc[0].Team} "
              f"{results[label].iloc[0].xPts:.1f}")

    # sheets
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    xlsx = OUT / "nhl_2026_27_projections_v2.xlsx"
    ledger = pd.read_csv(OUT / "feature_ledger.csv")
    ledger["status"] = "in model (R3)"
    extra = pd.DataFrame([
        {"feature": "heteroscedastic goalie sigma", "status":
         f"dropped (prereg coverage test; spread {p2['hetero']['spread']})"},
        {"feature": "roster-delta overlay (July 2026)", "status":
         f"mechanism rho={p2['prereg']['mechanism']['rho_hat']}, placebo "
         f"{p2['prereg']['mechanism']['placebo']} -> BLOCKED by prereg rule"},
        {"feature": "trajectory/momentum (v1 test)", "status": "null (gamma=0 optimal), dropped"},
    ])
    ledger_full = pd.concat([ledger, extra], ignore_index=True)
    movers = pd.read_csv(OUT / "roster_delta_2026.csv")
    a_struct, curves = age_sheet(fb, sk, bios)

    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        pd.DataFrame({"NHL 2026-27 Model v2": readme_lines(p2, aggs["2026_27"])}).to_excel(
            xw, sheet_name="README", index=False)
        results["2026_27"].to_excel(xw, sheet_name="Projections_2026_27", index=False)
        results["2027_28"].to_excel(xw, sheet_name="Projections_2027_28", index=False)
        goalie_sheet(fb, bios, go, ros).to_excel(xw, sheet_name="Goalies", index=False)
        fb.team_goalies(V_PROD).round(4).reset_index().to_excel(
            xw, sheet_name="Team_Goaltending", index=False)
        a_struct.to_excel(xw, sheet_name="Age_Structure", index=False, startrow=0)
        curves.to_excel(xw, sheet_name="Age_Structure", index=False,
                        startrow=len(a_struct) + 3)
        player_points_sheet(fb, sk, bios, ros).to_excel(
            xw, sheet_name="Player_Points", index=False)
        movers.to_excel(xw, sheet_name="Roster_Delta_2026", index=False)
        ledger_full.to_excel(xw, sheet_name="Feature_Ledger", index=False)
        backtest_sheet(p2).to_excel(xw, sheet_name="Backtest_v2", index=False)

        wb = xw.book
        pct_cols = {"Playoff%", "Division%", "Presidents%", "Conference%", "Cup%", "P(>=100pts)"}
        num1 = {"xPts", "SD", "P5", "P25", "P50", "P75", "P95", "Ridge dev (pts/82)",
                "Overlay dElo", "Sigma (Elo)"} | set(BLOCKS)
        hdr_fill = PatternFill("solid", fgColor="0B3D2E")
        hdr_font = Font(bold=True, color="FFFFFF")
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.fill = hdr_fill
                cell.font = hdr_font
                cell.alignment = Alignment(vertical="center")
            for j, hcell in enumerate(ws[1], start=1):
                hname = hcell.value
                width = 15
                if hname in ("Team", "Player", "Name", "NHL 2026-27 Model v2", "detail"):
                    width = 100 if ws.title in ("README",) else (60 if hname == "detail" else 26)
                ws.column_dimensions[get_column_letter(j)].width = width
                if hname in pct_cols:
                    for row_ in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                        row_[0].number_format = "0.0%"
                elif hname in num1:
                    for row_ in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                        row_[0].number_format = "0.0"
    results["2026_27"].to_csv(OUT / "projections_2026_27_v2.csv", index=False)
    results["2027_28"].to_csv(OUT / "projections_2027_28_v2.csv", index=False)

    # v1 preservation check
    chk = subprocess.run(["shasum", "-c", "SHA1SUMS"], cwd=OUT / "v1",
                         capture_output=True, text=True)
    assert chk.returncode == 0, f"v1 outputs modified!\n{chk.stdout}{chk.stderr}"
    print("v1 snapshot verified byte-identical")
    print(f"wrote {xlsx}")


if __name__ == "__main__":
    main()
