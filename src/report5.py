"""v5 production report (PLAN_V5): expanded-data features, gated in by F3.

v5 = v4 pipeline with the F3 gate survivors added to the ridge (h1: corsi_dev,
fo_dev, pen_diff per params_v5; h2: per params_v5). Everything else (overlay,
availability 2.0, goalie layer, sigma, b2b, n0_g) inherits from params_v4 shipped.

Per PLAN_V5 H: the live 2026-27 holdout still scores v1/v4/HOWE — v5's 2026-27
numbers are REPORT-ONLY context; v5 and HOWE5 = 0.5*v1 + 0.5*v5 are the 2027-28
co-headline. Seeds: 511 (h1) / 522 (h2), offsets v5=0 / HOWE5=1 — fresh stream,
deterministic. Writes output/nhl_projections_v5.xlsx + CSVs + v5_prior_ratings.csv.
v1-v4 outputs untouched.
"""
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import engine as E
import goalie_game as GG
import players as P
import ridge as R
from features_v5 import FEATURES_V5_ALL, FeatureBuilderV5
from overlay import points_per_goal, skater_value_goals   # noqa: F401 (parity w/ v4)
from report import FULL
from report4 import BLOCKS_V4, prod_noise, real_schedule_flags, summarize

PROJ = Path(__file__).resolve().parents[1]
RAW = PROJ / "data" / "raw"
OUT = PROJ / "output"
N_SIMS = 10_000
V_PROD = 2026

BLOCKS_V5 = dict(BLOCKS_V4)
BLOCKS_V5["Contrib NewData"] = ["corsi_dev", "fo_dev", "pen_diff", "hd_share",
                                "flurry_xg_dev", "rush_xg_dev"]


def production_predictions_v5(fb, feats, lam, h, target_T):
    pred_seasons = list(range(2012 + (h - 1), 2027)) + [target_T]
    X, y, meta = fb.feature_matrix(pred_seasons, h, feats=FEATURES_V5_ALL)
    idx = [FEATURES_V5_ALL.index(f) for f in feats]
    Xs = X[:, idx]
    tr = ~np.isnan(y)
    b = R.fit_ridge(Xs[tr], y[tr], lam)
    te = meta["T"].to_numpy() == target_T
    yhat = Xs[te] @ b
    yhat -= yhat.mean()
    teams = meta.loc[te, "team"].to_numpy()
    contribs = {}
    for label, fs in BLOCKS_V5.items():
        sub = [f for f in fs if f in feats]
        if not sub:
            continue
        ii = [feats.index(f) for f in sub]
        cv = Xs[te][:, ii] @ b[ii]
        contribs[label] = pd.Series(cv - cv.mean(), index=teams)
    return pd.Series(yhat, index=teams), contribs


def main():
    p2 = json.loads((OUT / "params_v2.json").read_text())
    p4 = json.loads((OUT / "params_v4.json").read_text())
    p5 = json.loads((OUT / "params_v5.json").read_text())
    v1p = json.loads((OUT / "params.json").read_text())
    ship4, ship5 = p4["shipped"], p5["shipped"]
    c, k = p2["c"], p2["k"]
    g, ts = E.load()
    preds, end_r, _ = E.run_elo(g, K=v1p["K"], H=v1p["H"], phi_s=v1p["phi_s"])
    fb = FeatureBuilderV5(end_r, ts, goalie_hp=p2["goalie_hp"],
                          skater_delta=p2["skater_delta"])
    om = E.fit_outcome(preds, list(range(2006, 2027)))
    beta_prod = E.fit_xg_beta(end_r, ts, list(range(2008, 2027)))
    sk, go, skt, got, bios = P.load_panels()
    delo = pd.read_csv(OUT / "overlay_team_deltas_v4.csv", index_col=0)["dElo"]
    sigma = ship4["sigma_c4"]

    results, prior_rows = {}, {}
    for h, (label, target_T, seed) in {1: ("2026_27", 2027, 511),
                                       2: ("2027_28", 2028, 522)}.items():
        feats = ship5[f"features_h{h}"]
        lam = ship5[f"lam_h{h}"]
        yhat, contribs = production_predictions_v5(fb, feats, lam, h, target_T)
        ov_scale = 1.0 if h == 1 else 0.88
        r_v5 = {t: 1505.0 + yhat[t] / c + ov_scale * float(delo.get(t, 0.0))
                for t in yhat.index}
        phi = v1p["phi1"] if h == 1 else v1p["phi2"]
        r_v1 = E.project_ratings(end_r, ts, beta_prod, V_PROD, w=v1p["w"], phi=phi)
        r_howe5 = {t: 1505.0 + 0.5 * (r_v1[t] - 1505.0) + 0.5 * (r_v5[t] - 1505.0)
                   for t in r_v5}

        noise_fn, nteams, par2, gpar, tand, frag = prod_noise(
            fb, sk, got, bios, ts, k, h, ship4["n0_g"])
        rng = np.random.default_rng(seed)
        if h == 1:
            sched = real_schedule_flags()
        else:
            sched = E.synthetic_schedule_84(E.DIVISIONS_CURRENT, rng)
            sched["hb2b"] = 0.0
            sched["ab2b"] = 0.0
            sched["d_adj"] = 0.0
        gn = (GG.make_game_noise(sched, tand, k) if ship4["goalie_layer"] else None)
        sim_cols = ["home", "away", "d_adj"]
        for name, ratings, seed_off in (("v5", r_v5, 0), ("HOWE5", r_howe5, 1)):
            assert sorted(ratings) == sorted(nteams)
            sim = E.simulate_season(ratings, sigma, sched[sim_cols], om,
                                    E.DIVISIONS_CURRENT, N_SIMS,
                                    np.random.default_rng(seed + seed_off),
                                    playoffs=True, extra_noise=noise_fn, game_noise=gn)
            assert sim["made_po"].sum(1).std() == 0 and sim["won_cup"].sum(1).mean() == 1.0
            extra = {"Avail SD (Elo)": frag.avail_sd_elo,
                     "Overlay dElo": pd.Series({t: round(ov_scale * float(delo.get(t, 0.0)), 1)
                                                for t in ratings})}
            if name == "v5":
                for lbl in contribs:
                    extra[lbl] = contribs[lbl]
            results[(label, name)] = summarize(sim, extra)
            print(f"{label} {name}: league mean {sim['pts'].mean():.1f}; top "
                  f"{results[(label, name)].iloc[0].Team} "
                  f"{results[(label, name)].iloc[0].xPts:.1f}")
        if h == 1:
            prior_rows = {t: {"rating_v1": round(r_v1[t], 2),
                              "rating_v5": round(r_v5[t], 2),
                              "rating_howe5": round(r_howe5[t], 2)}
                          for t in sorted(r_v5)}

    pd.DataFrame(prior_rows).T.rename_axis("team").to_csv(OUT / "v5_prior_ratings.csv")

    readme = [
        "NHL 2026-27 / 2027-28 — MODEL v5 (built 2026-08-19; PLAN_V5 prereg)",
        "",
        "v5 = v4 + expanded-data features gated by F3 (train-only, 2012-2017):",
        f"h1 adds {p5['gate_F3']['h1']['joint']['final_set_new']} "
        f"(train LOSO {p5['gate_F3']['h1']['mae_incumbent']} -> "
        f"{p5['gate_F3']['h1']['joint']['mae_final']}),",
        f"h2 adds {p5['gate_F3']['h2']['joint']['final_set_new']} "
        f"(train LOSO {p5['gate_F3']['h2']['mae_incumbent']} -> "
        f"{p5['gate_F3']['h2']['joint']['mae_final']}).",
        "Data: MoneyPuck shot-level, NHL play-by-play (2012-2026, cross-checked to "
        "r>0.99 vs official aggregates), NHL stats-rest reports, hockey-reference "
        "SRS/SOS, NHL EDGE (report-only), fastRhockey bulk (cross-check).",
        "",
        "IMPORTANT (PLAN_V5 H): the pre-registered live 2026-27 holdout still scores "
        "v1/v4/HOWE. The 2026-27 v5 sheets here are REPORT-ONLY context. v5 and "
        "HOWE5 = 0.5*v1 + 0.5*v5 are the 2027-28 co-headline.",
    ]
    xlsx = OUT / "nhl_projections_v5.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        pd.DataFrame({"NHL Model v5": readme}).to_excel(xw, sheet_name="README", index=False)
        results[("2027_28", "HOWE5")].to_excel(xw, sheet_name="Projections_2027_28_HOWE5", index=False)
        results[("2027_28", "v5")].to_excel(xw, sheet_name="Projections_2027_28_v5", index=False)
        results[("2026_27", "HOWE5")].to_excel(xw, sheet_name="ReportOnly_2026_27_HOWE5", index=False)
        results[("2026_27", "v5")].to_excel(xw, sheet_name="ReportOnly_2026_27_v5", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="1E3A5F")
                cell.font = Font(bold=True, color="FFFFFF")
    results[("2026_27", "v5")].to_csv(OUT / "projections_2026_27_v5.csv", index=False)
    results[("2027_28", "v5")].to_csv(OUT / "projections_2027_28_v5.csv", index=False)
    results[("2026_27", "HOWE5")].to_csv(OUT / "projections_2026_27_howe5.csv", index=False)
    results[("2027_28", "HOWE5")].to_csv(OUT / "projections_2027_28_howe5.csv", index=False)
    print(f"wrote {xlsx}")


if __name__ == "__main__":
    main()
