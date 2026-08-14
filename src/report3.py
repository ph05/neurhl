"""v3 production report: 18-feature ridge + availability noise in sims + living-model sheet.
Writes output/nhl_2026_27_projections_v3.xlsx, v3_prior_ratings.csv (consumed by live.py),
CSVs. v1/v2 outputs untouched."""
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import availability as A
import engine as E
import players as P
import ridge as R
from features import FEATURES_V3, FeatureBuilder
from overlay import points_per_goal, skater_value_goals
from report import FULL, fair_odds
from report2 import BLOCKS, real_schedule

PROJ = Path(__file__).resolve().parents[1]
RAW = PROJ / "data" / "raw"
OUT = PROJ / "output"
N_SIMS = 10_000
V_PROD = 2026

BLOCKS_V3 = dict(BLOCKS)
BLOCKS_V3["Contrib Finishing/PP"] = ["finishing", "st_pp"]


def production_predictions(fb, feats, lam, h, target_T):
    pred_seasons = list(range(2012 + (h - 1), 2027)) + [target_T]
    X, y, meta = fb.feature_matrix(pred_seasons, h, feats=FEATURES_V3)
    idx = [FEATURES_V3.index(f) for f in feats]
    Xs = X[:, idx]
    seasons = meta["T"].to_numpy()
    tr = ~np.isnan(y)
    b = R.fit_ridge(Xs[tr], y[tr], lam)
    te = seasons == target_T
    yhat = Xs[te] @ b
    yhat -= yhat.mean()
    teams = meta.loc[te, "team"].to_numpy()
    contribs = {}
    for label, fs in BLOCKS_V3.items():
        sub = [f for f in fs if f in feats]
        if not sub:
            continue
        ii = [feats.index(f) for f in sub]
        cv = Xs[te][:, ii] @ b[ii]
        contribs[label] = pd.Series(cv - cv.mean(), index=teams)
    return pd.Series(yhat, index=teams), contribs


def prod_availability_noise(fb, sk, skt, bios, ts, k):
    """Availability closure for the 2026-27 sim using August API rosters."""
    par = A.fit_availability(sk, bios, V_PROD)
    marcel = fb.marcel(V_PROD, 1)
    repl = P.replacement_rates(sk, V_PROD)
    ppg = points_per_goal(sk, ts, V_PROD)
    vals = skater_value_goals(marcel, repl, ppg)
    ros = pd.read_csv(RAW / "nhl_rosters_20262027.csv")
    ros = ros[ros.position != "G"]
    rosters = {}
    frag_rows = []
    for team, d in ros.groupby("team"):
        vv = vals.reindex(d.playerId).fillna(0.0)
        aa = P.age_of(bios, d.playerId, 2027)
        top = sorted(zip(aa.to_numpy(), vv.to_numpy()), key=lambda t_: -t_[1])[:A.TOP_N]
        rosters[team] = [(float(a_) if np.isfinite(a_) else 27.0, float(v_))
                        for a_, v_ in top]
    teams = sorted(rosters)
    fn = A.make_extra_noise(teams, rosters, par, k)
    probe = fn(4000, np.random.default_rng(2))
    for j, t in enumerate(teams):
        frag_rows.append({"team": t, "avail_sd_elo": float(probe[:, j].std())})
    return fn, teams, par, pd.DataFrame(frag_rows).set_index("team")


def main():
    p2 = json.loads((OUT / "params_v2.json").read_text())
    p3 = json.loads((OUT / "params_v3.json").read_text())
    v1 = json.loads((OUT / "params.json").read_text())
    assert "living_model" in p3, "run live.py (tune/validate) first"
    feats, lam1, lam2 = p3["final_feature_set"], p3["lam_h1"], p3["lam_h2"]
    c, k = p2["c"], p2["k"]
    g, ts = E.load()
    preds, end_r, _ = E.run_elo(g, K=v1["K"], H=v1["H"], phi_s=v1["phi_s"])
    fb = FeatureBuilder(end_r, ts, goalie_hp=p2["goalie_hp"], skater_delta=p2["skater_delta"])
    om = E.fit_outcome(preds, list(range(2006, 2027)))
    sk, go, skt, got, bios = P.load_panels()

    noise_fn, noise_teams, avail_par, frag = prod_availability_noise(fb, sk, skt, bios, ts, k)
    use_noise = p3["availability"]["keep"]
    sigma = p3["availability"]["sigma_c3"] if use_noise else p2["sigma_c"]["h1"]

    results = {}
    for h, (label, target_T, seed, lam) in {1: ("2026_27", 2027, 311, lam1),
                                            2: ("2027_28", 2028, 322, lam2)}.items():
        yhat, contribs = production_predictions(fb, feats, lam, h, target_T)
        ratings = {t: 1505.0 + yhat[t] / c for t in yhat.index}
        rng = np.random.default_rng(seed)
        sched = real_schedule() if h == 1 else E.synthetic_schedule_84(E.DIVISIONS_CURRENT, rng)
        assert sorted(ratings) == sorted(noise_teams)
        sim = E.simulate_season(ratings, sigma, sched, om, E.DIVISIONS_CURRENT, N_SIMS, rng,
                                playoffs=True, extra_noise=noise_fn if use_noise else None)
        assert sim["made_po"].sum(1).std() == 0 and sim["won_cup"].sum(1).mean() == 1.0
        pts = sim["pts"].astype(float)
        rows = []
        div_of = {t: dv for dv, ts_ in E.DIVISIONS_CURRENT.items() for t in ts_}
        for i, t in enumerate(sim["teams"]):
            p_ = pts[:, i]
            playoff = sim["made_po"][:, i].mean()
            cup = sim["won_cup"][:, i].mean()
            cup_dec, cup_amer = fair_odds(cup)
            row = {"Team": FULL[t], "Abbr": t, "Division": div_of[t],
                   "xPts": p_.mean(), "SD": p_.std(ddof=1),
                   "P5": np.percentile(p_, 5), "P50": np.percentile(p_, 50),
                   "P95": np.percentile(p_, 95),
                   "Playoff%": playoff, "Division%": sim["won_div"][:, i].mean(),
                   "Conference%": sim["won_conf"][:, i].mean(), "Cup%": cup,
                   "Avail SD (Elo)": float(frag.avail_sd_elo.get(t, 0.0)),
                   "FairOdds Cup (US)": cup_amer}
            for lbl in contribs:
                row[lbl] = float(contribs[lbl].get(t, 0.0))
            rows.append(row)
        df = pd.DataFrame(rows).sort_values("xPts", ascending=False).reset_index(drop=True)
        df.insert(0, "Rank", np.arange(1, 33))
        results[label] = df
        print(f"{label}: league mean {pts.mean():.1f}; top {df.iloc[0].Team} {df.iloc[0].xPts:.1f}")
        if h == 1:
            pd.Series({t: ratings[t] for t in sorted(ratings)}, name="rating").to_csv(
                OUT / "v3_prior_ratings.csv")

    # sheets
    fin = P.finishing_project(sk, V_PROD)
    names = sk.drop_duplicates("playerId").set_index("playerId")["name"]
    fin["Name"] = fin.playerId.map(names)
    fin["G>xG per season (proj)"] = (fin.theta_fin * fin.sog82).round(1)
    fin["per 100 shots"] = (fin.theta_fin * 100).round(2)
    fin_sheet = fin[fin.sog_win >= 300].nlargest(120, "theta_fin")[
        ["Name", "per 100 shots", "G>xG per season (proj)", "sog82", "sog_win"]].rename(
        columns={"sog82": "Proj SOG/82", "sog_win": "SOG (window)"})

    avail_sheet = pd.DataFrame([
        {"Age bucket": f"{lo}-{hi}", "Beta a": round(p_["a"], 2), "Beta b": round(p_["b"], 2),
         "E[GP share]": round(p_["mean"], 3), "n (fit pairs)": p_["n"]}
        for (lo, hi), p_ in zip(A.BUCKETS, avail_par)])
    frag_sheet = frag.round(1).reset_index()
    frag_sheet["Team"] = frag_sheet.team.map(FULL)

    lm = p3["living_model"]
    lm_rows = [{"item": "n0 (prior worth in games)", "value": lm["n0"]}]
    lm_rows += [{"item": f"tune rest-MAE n0={k_}", "value": v_}
                for k_, v_ in lm["tune_table"].items()]
    lm_rows += [{"item": f"VALID gp={k_} {c_}", "value": v_}
                for k_, r_ in lm["validation_mae"].items() for c_, v_ in r_.items()]
    lm_rows += [{"item": f"VALID playoff Brier gp={k_}", "value": v_}
                for k_, v_ in lm["brier_by_checkpoint"].items()]
    lm_rows.append({"item": "operate", "value": "python src/live.py update (nightly from Sep 29)"})

    led2 = pd.read_csv(OUT / "feature_ledger.csv")
    led3 = pd.DataFrame(p3["ledger_v3"])
    led3["source"] = "v3 gate"
    ledger = pd.concat([led2.assign(source="v2 stage_t"), led3], ignore_index=True)

    xlsx = OUT / "nhl_2026_27_projections_v3.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        readme = build_readme(p3, results)
        pd.DataFrame({"NHL 2026-27 Model v3": readme}).to_excel(xw, sheet_name="README", index=False)
        results["2026_27"].to_excel(xw, sheet_name="Projections_2026_27", index=False)
        results["2027_28"].to_excel(xw, sheet_name="Projections_2027_28", index=False)
        fin_sheet.to_excel(xw, sheet_name="Finishing", index=False)
        avail_sheet.to_excel(xw, sheet_name="Availability_Model", index=False)
        frag_sheet[["Team", "avail_sd_elo"]].rename(
            columns={"avail_sd_elo": "Availability SD (Elo)"}).to_excel(
            xw, sheet_name="Team_Fragility", index=False)
        pd.DataFrame(lm_rows).to_excel(xw, sheet_name="Living_Model", index=False)
        ledger.to_excel(xw, sheet_name="Feature_Ledger_v3", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="3B0764")
                cell.font = Font(bold=True, color="FFFFFF")
    results["2026_27"].to_csv(OUT / "projections_2026_27_v3.csv", index=False)
    print(f"wrote {xlsx}")


def build_readme(p3, results):
    L = []
    a = L.append
    a("MODEL v3 - the three commissioned enhancements, each pre-registered and gated on train only.")
    a("")
    a("1. FINISHING SKILL: EB per-shot (G - ixG); YoY r=0.31, stabilizes ~510 shots. GATE: KEPT")
    a(f"   (train LOSO dMAE {p3['ledger_v3'][0]['add_one_in_dMAE']:+.3f}, sign-stable 5/5 folds).")
    a("2. AVAILABILITY: beta-binomial games-played by age (24x overdispersed vs binomial;")
    a("   P(miss 15+) 29%->42% young->old). Enters sims as ZERO-MEAN per-player draws for each")
    a("   team's top-9 (mean effect stays with the age features; this term adds variance/tails).")
    a(f"   GATE: {'KEPT' if p3['availability']['keep'] else 'DROPPED'} "
      f"(coverage {p3['availability']['coverage']}).")
    a("3. LIVING MODEL: in-season posterior w = n0/(n0+games). Empirics: at 20 games, early pace")
    a("   and preseason knowledge carry equal predictive weight (betas .31/.31) => prior worth")
    a(f"   ~20 games; tuned n0 = {p3['living_model']['n0']}. Validated on 2022-2026 replays (see")
    a("   Living_Model sheet: blend vs pure-prior vs pure-Elo, playoff Brier by checkpoint).")
    a("   Also gated this round: PP process KEPT, PK process DROPPED (+0.05 MAE, documented).")
    a("")
    a("Projection sheets include per-block attribution incl. the new Finishing/PP column.")
    a("The 2018-2026 preseason windows were NOT touched by any v3 decision; the live 2026-27")
    a("season is the out-of-sample test for everything above. Operate: python src/live.py update.")
    return L


if __name__ == "__main__":
    main()
