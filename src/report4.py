"""v4 production report: gated improvements + locked HOWE co-headline (PLAN_V4).

HOWE = Hockey Outcomes via Weighted Ensemble (prereg name "ENS"; renamed 2026-08-19,
math/seeds unchanged — see src/howe.py).

Ships: 19-feature ridge (prospect_pipeline in at h1+h2), amended overlay (rho*=0.421),
availability 2.0 (bucket + zero-GP mass + goalie slot; h2 uses h2 ages), goalie
game layer (starter rotation, b2b backup rule), b2b d_adj on the real schedule,
float64 tiebreakers, HOWE = 0.5*v1 + 0.5*v4 (decision locked before computation).
Writes output/nhl_2026_27_projections_v4.xlsx + CSVs + v4_prior_ratings.csv
(v1/v4/HOWE columns, consumed by live.py). v1/v2/v3 outputs untouched.
"""
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import availability2 as A2
import engine as E
import goalie_game as GG
import players as P
import ridge as R
import scoring as SC
from features import FEATURES_V4, FeatureBuilder
from overlay import points_per_goal, skater_value_goals
from report import FULL, fair_odds
from report2 import BLOCKS

PROJ = Path(__file__).resolve().parents[1]
RAW = PROJ / "data" / "raw"
OUT = PROJ / "output"
N_SIMS = 10_000
V_PROD = 2026
B2B_ELO = 38.0

BLOCKS_V4 = dict(BLOCKS)
BLOCKS_V4["Contrib Finishing/PP"] = ["finishing", "st_pp"]
BLOCKS_V4["Contrib Prospects"] = ["prospect_pipeline"]


def real_schedule_flags() -> pd.DataFrame:
    df = pd.read_csv(RAW / "nhl_schedule_20262027.csv", parse_dates=["date"]) \
        .sort_values(["date", "game_id"]).reset_index(drop=True)
    hb, ab = E.b2b_flags(df[["date", "home", "away"]])
    df["hb2b"], df["ab2b"] = hb, ab
    df["d_adj"] = B2B_ELO * (df.ab2b - df.hb2b)
    return df[["date", "home", "away", "hb2b", "ab2b", "d_adj"]]


def production_predictions(fb, feats, lam, h, target_T):
    pred_seasons = list(range(2012 + (h - 1), 2027)) + [target_T]
    X, y, meta = fb.feature_matrix(pred_seasons, h, feats=FEATURES_V4)
    idx = [FEATURES_V4.index(f) for f in feats]
    Xs = X[:, idx]
    tr = ~np.isnan(y)
    b = R.fit_ridge(Xs[tr], y[tr], lam)
    te = meta["T"].to_numpy() == target_T
    yhat = Xs[te] @ b
    yhat -= yhat.mean()
    teams = meta.loc[te, "team"].to_numpy()
    contribs = {}
    for label, fs in BLOCKS_V4.items():
        sub = [f for f in fs if f in feats]
        if not sub:
            continue
        ii = [feats.index(f) for f in sub]
        cv = Xs[te][:, ii] @ b[ii]
        contribs[label] = pd.Series(cv - cv.mean(), index=teams)
    return pd.Series(yhat, index=teams), contribs


def prod_noise(fb, sk, got, bios, ts, k, h: int, n0_g: float):
    """Availability-2.0 closure + goalie game inputs from August rosters.
    h2 uses h2 ages (PLAN_V4 B6)."""
    age_ref = 2027 + (h - 1)
    par2 = A2.fit_availability2(sk, bios, V_PROD)
    gpar = A2.fit_goalie_availability(got, ts, V_PROD)
    marcel = fb.marcel(V_PROD, h)
    repl = P.replacement_rates(sk, V_PROD)
    ppg = points_per_goal(sk, ts, V_PROD)
    vals = skater_value_goals(marcel, repl, ppg)
    ros = pd.read_csv(RAW / "nhl_rosters_20262027.csv")
    tand = GG.prod_tandems(ros, got, ts, fb.goalie_proj(V_PROD), V_PROD, n0_g)
    rosters2, frag_rows = {}, []
    for team, d in ros[ros.position != "G"].groupby("team"):
        vv = vals.reindex(d.playerId).fillna(0.0)
        aa = P.age_of(bios, d.playerId, age_ref)
        top = sorted(zip(aa.to_numpy(), vv.to_numpy()), key=lambda t_: -t_[1])[:A2.TOP_N]
        b_idx = A2._bucket_idx(np.array([a if np.isfinite(a) else 27.0 for a, _ in top]))
        rosters2[team] = [(int(bi), par2["buckets"][int(bi)]["mean"], float(v_))
                          for bi, (_, v_) in zip(b_idx, top)]
    grow = {t: (gpar["mean"], max(float(tand.loc[t].theta1 - tand.loc[t].theta2), 0.0)
                * A2.GOALIE_SHOTS) for t in tand.index}
    teams = sorted(rosters2)
    fn = A2.make_extra_noise2(teams, rosters2, par2, grow, gpar, k)
    probe = fn(4000, np.random.default_rng(2))
    for j, t in enumerate(teams):
        frag_rows.append({"team": t, "avail_sd_elo": float(probe[:, j].std())})
    return fn, teams, par2, gpar, tand, pd.DataFrame(frag_rows).set_index("team")


def summarize(sim, extra_cols: dict) -> pd.DataFrame:
    pts = sim["pts"].astype(float)
    key = E.standings_key(pts, sim["rw"].astype(float), sim["row"].astype(float),
                          sim["win"].astype(float),
                          np.random.default_rng(1).random(pts.shape))
    top_idx = key.argmax(axis=1)
    div_of = {t: dv for dv, ts_ in E.DIVISIONS_CURRENT.items() for t in ts_}
    rows = []
    for i, t in enumerate(sim["teams"]):
        p_ = pts[:, i]
        cup = sim["won_cup"][:, i].mean()
        _, cup_amer = fair_odds(cup)
        row = {"Team": FULL[t], "Abbr": t, "Division": div_of[t],
               "xPts": p_.mean(), "SD": p_.std(ddof=1),
               "P5": np.percentile(p_, 5), "P50": np.percentile(p_, 50),
               "P95": np.percentile(p_, 95),
               "Playoff%": sim["made_po"][:, i].mean(),
               "Division%": sim["won_div"][:, i].mean(),
               "Presidents%": float((top_idx == i).mean()),
               "Conference%": sim["won_conf"][:, i].mean(), "Cup%": cup,
               "FairOdds Cup (US)": cup_amer}
        for cname, series in extra_cols.items():
            row[cname] = float(series.get(t, 0.0)) if hasattr(series, "get") else series
        rows.append(row)
    df = pd.DataFrame(rows).sort_values("xPts", ascending=False).reset_index(drop=True)
    df.insert(0, "Rank", np.arange(1, len(df) + 1))
    return df


def main():
    p2 = json.loads((OUT / "params_v2.json").read_text())
    p4 = json.loads((OUT / "params_v4.json").read_text())
    v1p = json.loads((OUT / "params.json").read_text())
    ship = p4["shipped"]
    c, k = p2["c"], p2["k"]
    g, ts = E.load()
    preds, end_r, _ = E.run_elo(g, K=v1p["K"], H=v1p["H"], phi_s=v1p["phi_s"])
    fb = FeatureBuilder(end_r, ts, goalie_hp=p2["goalie_hp"],
                        skater_delta=p2["skater_delta"])
    om = E.fit_outcome(preds, list(range(2006, 2027)))
    beta_prod = E.fit_xg_beta(end_r, ts, list(range(2008, 2027)))
    sk, go, skt, got, bios = P.load_panels()
    delo = pd.read_csv(OUT / "overlay_team_deltas_v4.csv", index_col=0)["dElo"]
    sigma = ship["sigma_c4"]

    results, prior_rows = {}, {}
    for h, (label, target_T, seed) in {1: ("2026_27", 2027, 411),
                                       2: ("2027_28", 2028, 422)}.items():
        feats = ship[f"features_h{h}"]
        lam = ship[f"lam_h{h}"]
        yhat, contribs = production_predictions(fb, feats, lam, h, target_T)
        ov_scale = 1.0 if h == 1 else 0.88
        r_v4 = {t: 1505.0 + yhat[t] / c + ov_scale * float(delo.get(t, 0.0))
                for t in yhat.index}
        phi = v1p["phi1"] if h == 1 else v1p["phi2"]
        r_v1 = E.project_ratings(end_r, ts, beta_prod, V_PROD, w=v1p["w"], phi=phi)
        r_howe = {t: 1505.0 + 0.5 * (r_v1[t] - 1505.0) + 0.5 * (r_v4[t] - 1505.0)
                  for t in r_v4}

        noise_fn, nteams, par2, gpar, tand, frag = prod_noise(
            fb, sk, got, bios, ts, k, h, ship["n0_g"])
        rng = np.random.default_rng(seed)
        if h == 1:
            sched = real_schedule_flags()
        else:
            sched = E.synthetic_schedule_84(E.DIVISIONS_CURRENT, rng)
            sched["hb2b"] = 0.0
            sched["ab2b"] = 0.0
            sched["d_adj"] = 0.0
        gn = (GG.make_game_noise(sched, tand, k) if ship["goalie_layer"] else None)
        sim_cols = ["home", "away", "d_adj"]
        for name, ratings, seed_off in (("v4", r_v4, 0), ("HOWE", r_howe, 1)):
            assert sorted(ratings) == sorted(nteams)
            sim = E.simulate_season(ratings, sigma, sched[sim_cols], om,
                                    E.DIVISIONS_CURRENT, N_SIMS,
                                    np.random.default_rng(seed + seed_off),
                                    playoffs=True, extra_noise=noise_fn, game_noise=gn)
            assert sim["made_po"].sum(1).std() == 0 and sim["won_cup"].sum(1).mean() == 1.0
            extra = {"Avail SD (Elo)": frag.avail_sd_elo,
                     "Overlay dElo": pd.Series({t: round(ov_scale * float(delo.get(t, 0.0)), 1)
                                                for t in ratings})}
            if name == "v4":
                for lbl in contribs:
                    extra[lbl] = contribs[lbl]
            results[(label, name)] = summarize(sim, extra)
            print(f"{label} {name}: league mean {sim['pts'].mean():.1f}; top "
                  f"{results[(label, name)].iloc[0].Team} "
                  f"{results[(label, name)].iloc[0].xPts:.1f}")
        if h == 1:
            prior_rows = {t: {"rating_v1": round(r_v1[t], 2),
                              "rating_v4": round(r_v4[t], 2),
                              "rating_howe": round(r_howe[t], 2)} for t in sorted(r_v4)}
            frag_h1, tand_h1, par2_h1, gpar_h1 = frag, tand, par2, gpar

    pd.DataFrame(prior_rows).T.rename_axis("team").to_csv(OUT / "v4_prior_ratings.csv")

    # ---- market sheet (hand-recorded board, 2026-08-17)
    board = SC.load_cup_board()
    market = None
    if board is not None:
        cup_v4 = results[("2026_27", "v4")].set_index("Abbr")["Cup%"]
        cup_howe = results[("2026_27", "HOWE")].set_index("Abbr")["Cup%"]
        market = SC.cup_market_sheet(cup_howe, board)
        market["model_v4"] = market.team.map(cup_v4)
        market = market.rename(columns={"model_p": "model_howe"})
        print(f"market board: overround {market.attrs['overround']:.3f}; "
              f"+EV (HOWE, both devigs): {list(market[market.stake > 0].team)}")

    # ---- sheets
    prospect_sheet = pd.DataFrame({
        "h1 (2026-27) pts": fb.team_prospects(V_PROD, 1, sorted(prior_rows)),
        "h2 (2027-28) pts": fb.team_prospects(V_PROD, 2, sorted(prior_rows)),
    }).rename_axis("team").round(1).reset_index()
    tand_sheet = tand_h1.round(4).reset_index()
    tand_sheet["Team"] = tand_sheet.team.map(FULL)
    avail_sheet = pd.DataFrame([
        {"Age bucket": f"{lo}-{hi}", "E[GP share]": round(p_["mean"], 3),
         "P(zero GP)": round(p_["p0"], 3), "Beta a": round(p_["a"], 2),
         "Beta b": round(p_["b"], 2), "n": p_["n"]}
        for (lo, hi), p_ in zip(A2.BUCKETS, par2_h1["buckets"])]
        + [{"Age bucket": "G1 goalies (pool)", "E[GP share]": round(gpar_h1["mean"], 3),
            "P(zero GP)": round(gpar_h1["p0"], 3), "Beta a": round(gpar_h1["a"], 2),
            "Beta b": round(gpar_h1["b"], 2), "n": gpar_h1["n"]}])
    gates_sheet = pd.DataFrame([
        {"gate": "G (goalie layer)", "result": json.dumps(p4["gate_G"])},
        {"gate": "S1/S2 + persistence", "result": json.dumps(p4["screens"]["S1"])
         + f" persistence_kept={p4['screens']['persistence_kept']}"},
        {"gate": "A2 (availability 2.0)", "result": json.dumps(p4["gate_A2"])},
        {"gate": "F2 h1", "result": json.dumps(p4["gate_F2"]["h1"])},
        {"gate": "F2 h2", "result": json.dumps(p4["gate_F2"]["h2"])},
        {"gate": "shipped", "result": json.dumps({kk: vv for kk, vv in ship.items()
                                                  if not kk.startswith("features")})},
    ])

    readme = [
        "NHL 2026-27 / 2027-28 — MODEL v4 (built 2026-08-17; PLAN_V4 prereg @ git 4bb92d9)",
        "",
        "CO-HEADLINES: Projections_2026_27_HOWE (locked 50/50 v1+v4 ensemble) and _v4.",
        "HOWE = Hockey Outcomes via Weighted Ensemble (prereg name ENS; renamed "
        "2026-08-19, math/seeds unchanged).",
        "The live 2026-27 season scores v1, v4 and HOWE — pre-registered holdout.",
        "",
        "Gated in this round (train-only, 2018-2026 untouched):",
        f"1. GOALIE GAME LAYER: starter rotation per game, b2b backup rule "
        f"(beta_b2b=0.5 declared), zero-mean (verified {p4['gate_G']['max_mean_shift']} pts). "
        f"GATE G PASS (coverage {p4['gate_G']['coverage']}).",
        f"2. AVAILABILITY 2.0: + zero-GP mass (P0 up to 14.6% at 36+), goalie G1 slot, "
        f"h2 ages. GATE A2 PASS (coverage {p4['gate_A2']['coverage']}, spread "
        f"{p4['gate_A2']['spread']} vs bound {p4['gate_A2']['spread_v3']}+0.02 — "
        f"passes by 0.004, watch live). Persistence DROPPED (EB lost to bucket-only).",
        f"3. PROSPECT PIPELINE: draft->panel join (98% on picks 1-15), train ramps; "
        f"GATE F2: h1 dMAE {p4['gate_F2']['h1']['dmae']:+.4f}, h2 "
        f"{p4['gate_F2']['h2']['dmae']:+.4f} (largest single train gain since v2), "
        f"sign-stable 1.00 — ENTERS both horizons.",
        "4. OVERLAY AMENDED + REPRODUCIBLE: rho spec-robust 0.843 with production-"
        "consistent weights -> rho*=0.421 (v3.1's 1.011 reproduced; its gain was the "
        "intercept fixing a weighting artifact, not the selection control).",
        "5. BUGS: float64 tiebreak key restores ROW (+W added); b2b constant 38 "
        "reproduced from committed code (t -5.3/+6.7); live.py rebuilt to the "
        "validated estimator with full odds output.",
        "",
        "Market_vs_Model: hand-recorded Cup board 2026-08-17 (overround 1.255). "
        "Fair odds are break-even prices, NOT bet instructions; see the stake column "
        "for quarter-Kelly at actual prices — it is deliberately tiny.",
    ]

    xlsx = OUT / "nhl_2026_27_projections_v4.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        pd.DataFrame({"NHL Model v4": readme}).to_excel(xw, sheet_name="README", index=False)
        results[("2026_27", "HOWE")].to_excel(xw, sheet_name="Projections_2026_27_HOWE", index=False)
        results[("2026_27", "v4")].to_excel(xw, sheet_name="Projections_2026_27_v4", index=False)
        results[("2027_28", "HOWE")].to_excel(xw, sheet_name="Projections_2027_28_HOWE", index=False)
        results[("2027_28", "v4")].to_excel(xw, sheet_name="Projections_2027_28_v4", index=False)
        if market is not None:
            market.round(4).to_excel(xw, sheet_name="Market_vs_Model", index=False)
        prospect_sheet.to_excel(xw, sheet_name="Prospect_Pipeline", index=False)
        tand_sheet.to_excel(xw, sheet_name="Goalie_Rotation", index=False)
        avail_sheet.to_excel(xw, sheet_name="Availability_v4", index=False)
        gates_sheet.to_excel(xw, sheet_name="Gates_v4", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="14532D")
                cell.font = Font(bold=True, color="FFFFFF")
    results[("2026_27", "HOWE")].to_csv(OUT / "projections_2026_27_howe.csv", index=False)
    results[("2026_27", "v4")].to_csv(OUT / "projections_2026_27_v4.csv", index=False)
    results[("2027_28", "HOWE")].to_csv(OUT / "projections_2027_28_howe.csv", index=False)
    results[("2027_28", "v4")].to_csv(OUT / "projections_2027_28_v4.csv", index=False)
    if market is not None:
        market.round(4).to_csv(OUT / "market_vs_model_2026_27.csv", index=False)
    print(f"wrote {xlsx}")


if __name__ == "__main__":
    main()
